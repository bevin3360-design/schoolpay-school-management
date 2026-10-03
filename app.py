from datetime import datetime
from io import BytesIO
import os
import re
import sqlite3
import zipfile
import xml.etree.ElementTree as ET
from functools import wraps
from typing import Any, Dict, List
from werkzeug.utils import secure_filename
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from flask import Flask, abort, flash, has_request_context, redirect, render_template, request, send_file, session, url_for

app = Flask(__name__)
app.secret_key = "school-report-secret"
app.config["UPLOAD_FOLDER"] = os.path.join(app.root_path, "static", "uploads")
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024
DB_PATH = os.path.join(app.root_path, "school.db")
os.makedirs(app.config["UPLOAD_FOLDER"], exist_ok=True)


def get_db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db() -> None:
    with get_db() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS schools (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                motto TEXT,
                logo_url TEXT,
                signature_url TEXT,
                report_template TEXT DEFAULT 'default'
            );

            CREATE TABLE IF NOT EXISTS teachers (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                password TEXT NOT NULL,
                full_name TEXT NOT NULL,
                school_id INTEGER DEFAULT 1,
                role TEXT DEFAULT 'admin',
                is_super_admin INTEGER DEFAULT 0,
                FOREIGN KEY(school_id) REFERENCES schools(id)
            );

            CREATE TABLE IF NOT EXISTS students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                admission_number TEXT NOT NULL,
                payment_code TEXT,
                full_name TEXT NOT NULL,
                gender TEXT,
                grade TEXT NOT NULL,
                school_stage TEXT,
                class_name TEXT,
                dob TEXT,
                parent_name TEXT,
                parent_relationship TEXT DEFAULT 'Guardian',
                parent_contact TEXT,
                school_id INTEGER DEFAULT 1,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(school_id) REFERENCES schools(id)
            );

            CREATE TABLE IF NOT EXISTS subjects (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                short_name TEXT,
                grade_level TEXT,
                school_id INTEGER DEFAULT 1,
                FOREIGN KEY(school_id) REFERENCES schools(id)
            );

            CREATE TABLE IF NOT EXISTS exams (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                exam_type TEXT NOT NULL,
                term TEXT,
                grade_level TEXT,
                school_id INTEGER DEFAULT 1,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(school_id) REFERENCES schools(id)
            );

            CREATE TABLE IF NOT EXISTS marks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                exam_id INTEGER NOT NULL,
                subject_id INTEGER NOT NULL,
                marks_obtained REAL NOT NULL,
                out_of REAL DEFAULT 100,
                school_id INTEGER DEFAULT 1,
                FOREIGN KEY(student_id) REFERENCES students(id),
                FOREIGN KEY(exam_id) REFERENCES exams(id),
                FOREIGN KEY(subject_id) REFERENCES subjects(id),
                FOREIGN KEY(school_id) REFERENCES schools(id)
            );

            CREATE TABLE IF NOT EXISTS teacher_learning_areas (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                school_id INTEGER NOT NULL,
                teacher_id INTEGER NOT NULL,
                subject_id INTEGER NOT NULL,
                grade TEXT NOT NULL,
                UNIQUE(teacher_id, subject_id, grade),
                FOREIGN KEY(school_id) REFERENCES schools(id),
                FOREIGN KEY(teacher_id) REFERENCES teachers(id),
                FOREIGN KEY(subject_id) REFERENCES subjects(id)
            );

            CREATE TABLE IF NOT EXISTS fee_items (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                school_id INTEGER NOT NULL,
                grade TEXT NOT NULL,
                academic_year TEXT NOT NULL,
                term TEXT NOT NULL,
                name TEXT NOT NULL,
                amount INTEGER NOT NULL CHECK(amount > 0),
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(school_id, grade, academic_year, term, name),
                FOREIGN KEY(school_id) REFERENCES schools(id)
            );

            CREATE TABLE IF NOT EXISTS fee_charges (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                school_id INTEGER NOT NULL,
                student_id INTEGER NOT NULL,
                fee_item_id INTEGER NOT NULL,
                item_name TEXT NOT NULL,
                academic_year TEXT NOT NULL,
                term TEXT NOT NULL,
                amount INTEGER NOT NULL CHECK(amount > 0),
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(student_id, fee_item_id),
                FOREIGN KEY(school_id) REFERENCES schools(id),
                FOREIGN KEY(student_id) REFERENCES students(id),
                FOREIGN KEY(fee_item_id) REFERENCES fee_items(id)
            );

            CREATE TABLE IF NOT EXISTS payments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                school_id INTEGER NOT NULL,
                amount INTEGER NOT NULL CHECK(amount > 0),
                channel TEXT NOT NULL,
                transaction_code TEXT,
                payer_phone TEXT,
                reference TEXT NOT NULL,
                academic_year TEXT,
                term TEXT,
                student_id INTEGER,
                suggested_student_id INTEGER,
                status TEXT NOT NULL CHECK(status IN ('matched', 'pending')),
                received_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                verified_by INTEGER,
                verified_at TEXT,
                verification_note TEXT,
                FOREIGN KEY(school_id) REFERENCES schools(id),
                FOREIGN KEY(student_id) REFERENCES students(id),
                FOREIGN KEY(suggested_student_id) REFERENCES students(id),
                FOREIGN KEY(verified_by) REFERENCES teachers(id)
            );

            CREATE UNIQUE INDEX IF NOT EXISTS idx_payments_transaction_code
            ON payments(school_id, transaction_code)
            WHERE transaction_code IS NOT NULL AND transaction_code != '';

            CREATE TABLE IF NOT EXISTS receipts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                school_id INTEGER NOT NULL,
                payment_id INTEGER NOT NULL UNIQUE,
                receipt_number TEXT NOT NULL UNIQUE,
                previous_balance INTEGER NOT NULL DEFAULT 0,
                new_balance INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(school_id) REFERENCES schools(id),
                FOREIGN KEY(payment_id) REFERENCES payments(id)
            );

            CREATE TABLE IF NOT EXISTS audit_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                school_id INTEGER NOT NULL,
                actor_id INTEGER,
                action TEXT NOT NULL,
                entity_type TEXT NOT NULL,
                entity_id INTEGER NOT NULL,
                details TEXT NOT NULL,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(school_id) REFERENCES schools(id),
                FOREIGN KEY(actor_id) REFERENCES teachers(id)
            );
            """
        )

        def table_has_column(table_name: str, column_name: str) -> bool:
            columns = conn.execute(f"PRAGMA table_info({table_name})").fetchall()
            return any(column[1] == column_name for column in columns)

        if not table_has_column("schools", "signature_url"):
            conn.execute("ALTER TABLE schools ADD COLUMN signature_url TEXT")
        if not table_has_column("schools", "report_template"):
            conn.execute("ALTER TABLE schools ADD COLUMN report_template TEXT DEFAULT 'default'")
        if not table_has_column("teachers", "school_id"):
            conn.execute("ALTER TABLE teachers ADD COLUMN school_id INTEGER DEFAULT 1")
        if not table_has_column("teachers", "role"):
            conn.execute("ALTER TABLE teachers ADD COLUMN role TEXT DEFAULT 'admin'")
        if not table_has_column("teachers", "is_super_admin"):
            conn.execute("ALTER TABLE teachers ADD COLUMN is_super_admin INTEGER DEFAULT 0")
        if not table_has_column("students", "school_id"):
            conn.execute("ALTER TABLE students ADD COLUMN school_id INTEGER DEFAULT 1")
        if not table_has_column("students", "payment_code"):
            conn.execute("ALTER TABLE students ADD COLUMN payment_code TEXT")
        if not table_has_column("students", "school_stage"):
            conn.execute("ALTER TABLE students ADD COLUMN school_stage TEXT")
        if not table_has_column("students", "parent_relationship"):
            conn.execute("ALTER TABLE students ADD COLUMN parent_relationship TEXT DEFAULT 'Guardian'")
        if not table_has_column("payments", "academic_year"):
            conn.execute("ALTER TABLE payments ADD COLUMN academic_year TEXT")
        if not table_has_column("payments", "term"):
            conn.execute("ALTER TABLE payments ADD COLUMN term TEXT")
        if not table_has_column("receipts", "previous_balance"):
            conn.execute("ALTER TABLE receipts ADD COLUMN previous_balance INTEGER NOT NULL DEFAULT 0")
        if not table_has_column("receipts", "new_balance"):
            conn.execute("ALTER TABLE receipts ADD COLUMN new_balance INTEGER NOT NULL DEFAULT 0")
        if not table_has_column("subjects", "school_id"):
            conn.execute("ALTER TABLE subjects ADD COLUMN school_id INTEGER DEFAULT 1")
        if not table_has_column("exams", "school_id"):
            conn.execute("ALTER TABLE exams ADD COLUMN school_id INTEGER DEFAULT 1")
        if not table_has_column("marks", "school_id"):
            conn.execute("ALTER TABLE marks ADD COLUMN school_id INTEGER DEFAULT 1")

        conn.execute("UPDATE teachers SET school_id = 1 WHERE school_id IS NULL")
        conn.execute("UPDATE students SET school_id = 1 WHERE school_id IS NULL")
        conn.execute("UPDATE subjects SET school_id = 1 WHERE school_id IS NULL")
        conn.execute("UPDATE exams SET school_id = 1 WHERE school_id IS NULL")
        conn.execute("UPDATE marks SET school_id = 1 WHERE school_id IS NULL")
        conn.execute(
            """
            UPDATE students
            SET school_stage = CASE
                WHEN grade IN ('Playgroup', 'PP1', 'PP2') THEN 'Pre-primary'
                WHEN grade IN ('Grade 7', 'Grade 8', 'Grade 9') THEN 'Junior School'
                WHEN grade IN ('Grade 10', 'Grade 11', 'Grade 12') THEN 'Senior School'
                ELSE 'Primary'
            END
            WHERE school_stage IS NULL OR school_stage = ''
            """
        )

        conn.execute(
            "INSERT OR IGNORE INTO schools (id, name, motto, logo_url, signature_url, report_template) VALUES (1, ?, ?, ?, ?, ?)",
            (
                "Sharleen School",
                "Nurturing Excellence, Character and Purpose",
                "/static/sample-report-form.png",
                "",
                "default",
            ),
        )
        conn.execute(
            "INSERT OR IGNORE INTO schools (id, name, motto, logo_url, signature_url, report_template) VALUES (2, ?, ?, ?, ?, ?)",
            (
                "Bright Stars Academy",
                "Learning with Purpose",
                "/static/sample-report-form.png",
                "",
                "default",
            ),
        )
        conn.execute(
            "INSERT OR IGNORE INTO teachers (id, username, password, full_name, school_id, role, is_super_admin) VALUES (1, ?, ?, ?, ?, ?, ?)",
            ("teacher", "password", "Teacher Admin", 1, "super_admin", 1),
        )
        conn.execute(
            "INSERT OR IGNORE INTO teachers (id, username, password, full_name, school_id, role, is_super_admin) VALUES (2, ?, ?, ?, ?, ?, ?)",
            ("teacher2", "password", "Teacher Admin 2", 2, "admin", 0),
        )

        default_subjects = [
            ("English", "ENG", "All"),
            ("Kiswahili", "KIS", "All"),
            ("Mathematics", "MATH", "All"),
            ("Science and Technology", "SCT", "All"),
            ("Integrated Science", "SCI", "All"),
            ("Social Studies", "SST", "All"),
            ("Agriculture", "AGR", "All"),
            ("Home Science", "HSC", "All"),
            ("Business Studies", "BUS", "All"),
            ("Computer Science", "CS", "All"),
            ("Pre-Technical Studies", "PTS", "All"),
            ("Creative Arts", "CA", "All"),
            ("Life Skills Education", "LSE", "All"),
            ("Health Education", "HE", "All"),
            ("Physical and Health Education", "PHE", "All"),
            ("Religious Education", "RE", "All"),
            ("Indigenous Languages", "IL", "All"),
            ("Kenya Sign Language", "KSL", "All"),
        ]
        for subject in default_subjects:
            conn.execute(
                "INSERT OR IGNORE INTO subjects (name, short_name, grade_level, school_id) VALUES (?, ?, ?, ?)",
                (*subject, 1),
            )
            conn.execute(
                "INSERT OR IGNORE INTO subjects (name, short_name, grade_level, school_id) VALUES (?, ?, ?, ?)",
                (*subject, 2),
            )

        default_exams = [
            ("Opening Assessment", "Opener", "Term 1", "All"),
            ("Midterm Assessment", "Midterm", "Term 1", "All"),
            ("End Term Assessment", "End Term", "Term 1", "All"),
        ]
        for exam in default_exams:
            conn.execute(
                "INSERT OR IGNORE INTO exams (name, exam_type, term, grade_level, school_id) VALUES (?, ?, ?, ?, ?)",
                (*exam, 1),
            )
            conn.execute(
                "INSERT OR IGNORE INTO exams (name, exam_type, term, grade_level, school_id) VALUES (?, ?, ?, ?, ?)",
                (*exam, 2),
            )

        default_students = [
            (
                "ADM001",
                "Amina Muli",
                "Female",
                "Grade 7",
                "JSS 1",
                "2014-05-10",
                "Muli Mwangi",
                "0721000111",
            ),
            (
                "ADM002",
                "Daniel Otieno",
                "Male",
                "Grade 8",
                "JSS 2",
                "2013-03-22",
                "Grace Otieno",
                "0721000222",
            ),
            (
                "ADM003",
                "Faith Wanjiku",
                "Female",
                "Grade 9",
                "JSS 3",
                "2012-08-04",
                "Joseph Wanjiku",
                "0721000333",
            ),
            (
                "ADM004",
                "Brian Kariuki",
                "Male",
                "Grade 6",
                "Primary 6",
                "2015-01-18",
                "Jane Kariuki",
                "0721000444",
            ),
            (
                "ADM005",
                "Lilian Njeri",
                "Female",
                "Grade 5",
                "Primary 5",
                "2016-10-02",
                "Charles Njeri",
                "0721000555",
            ),
        ]
        for student in default_students:
            conn.execute(
                "INSERT OR IGNORE INTO students (admission_number, full_name, gender, grade, class_name, dob, parent_name, parent_contact, school_id) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (*student, 1),
            )
        conn.execute(
            """
            UPDATE students
            SET school_stage = CASE
                WHEN grade IN ('Playgroup', 'PP1', 'PP2') THEN 'Pre-primary'
                WHEN grade IN ('Grade 7', 'Grade 8', 'Grade 9') THEN 'Junior School'
                WHEN grade IN ('Grade 10', 'Grade 11', 'Grade 12') THEN 'Senior School'
                ELSE 'Primary'
            END
            WHERE school_stage IS NULL OR school_stage = ''
            """
        )

        students_without_payment_codes = conn.execute(
            "SELECT id, school_id FROM students WHERE payment_code IS NULL OR payment_code = ''"
        ).fetchall()
        for student in students_without_payment_codes:
            conn.execute(
                "UPDATE students SET payment_code = ? WHERE id = ?",
                (f"S{student['school_id']}-{student['id']:06d}", student["id"]),
            )
        conn.execute(
            "CREATE UNIQUE INDEX IF NOT EXISTS idx_students_school_payment_code "
            "ON students(school_id, payment_code)"
        )

        if conn.execute("SELECT COUNT(*) AS count FROM marks").fetchone()["count"] == 0:
            sample_students = conn.execute("SELECT id FROM students WHERE school_id = 1 ORDER BY id").fetchall()
            sample_subjects = conn.execute("SELECT id FROM subjects WHERE school_id = 1 ORDER BY id").fetchall()
            sample_exam = conn.execute("SELECT id FROM exams WHERE school_id = 1 ORDER BY id LIMIT 1").fetchone()
            if sample_students and sample_subjects and sample_exam:
                for student in sample_students:
                    for subject in sample_subjects[:5]:
                        marks = 72 + (student["id"] % 3) * 5 + (subject["id"] % 2) * 3
                        conn.execute(
                            "INSERT INTO marks (student_id, exam_id, subject_id, marks_obtained, out_of, school_id) VALUES (?, ?, ?, ?, ?, ?)",
                            (student["id"], sample_exam["id"], subject["id"], marks, 100, 1),
                        )


init_db()


def parse_ksh_amount(value: str | None) -> int | None:
    try:
        amount = int(value or "")
    except (TypeError, ValueError):
        return None
    return amount if amount > 0 else None


def normalize_phone(value: str | None) -> str:
    digits = re.sub(r"\D", "", value or "")
    if digits.startswith("254") and len(digits) == 12:
        return "0" + digits[3:]
    return digits


def stage_for_grade(grade: str) -> str:
    if grade in {"Playgroup", "PP1", "PP2"}:
        return "Pre-primary"
    if grade in {"Grade 1", "Grade 2", "Grade 3", "Grade 4", "Grade 5", "Grade 6"}:
        return "Primary"
    if grade in {"Grade 7", "Grade 8", "Grade 9"}:
        return "Junior School"
    if grade in {"Grade 10", "Grade 11", "Grade 12"}:
        return "Senior School"
    return ""


def valid_student_stage(grade: str, school_stage: str) -> bool:
    return school_stage == stage_for_grade(grade)


def add_audit_log(
    conn: sqlite3.Connection,
    school_id: int,
    actor_id: int | None,
    action: str,
    entity_type: str,
    entity_id: int,
    details: str,
) -> None:
    conn.execute(
        """
        INSERT INTO audit_logs (school_id, actor_id, action, entity_type, entity_id, details)
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        (school_id, actor_id, action, entity_type, entity_id, details),
    )


def create_payment_receipt(
    conn: sqlite3.Connection,
    payment_id: int,
    school_id: int,
    actor_id: int | None,
) -> int:
    receipt = conn.execute(
        "SELECT id FROM receipts WHERE payment_id = ?", (payment_id,)
    ).fetchone()
    if receipt:
        return int(receipt["id"])

    payment = conn.execute(
        "SELECT student_id, amount, academic_year, term FROM payments WHERE id = ? AND school_id = ? AND status = 'matched'",
        (payment_id, school_id),
    ).fetchone()
    if not payment:
        raise ValueError("Only a matched payment can receive a receipt.")
    charges = conn.execute(
        """
        SELECT COALESCE(SUM(amount), 0) FROM fee_charges
        WHERE student_id = ? AND school_id = ? AND academic_year = ? AND term = ?
        """,
        (payment["student_id"], school_id, payment["academic_year"], payment["term"]),
    ).fetchone()[0]
    other_payments = conn.execute(
        """
        SELECT COALESCE(SUM(amount), 0) FROM payments
        WHERE student_id = ? AND school_id = ? AND status = 'matched'
          AND academic_year = ? AND term = ? AND id != ?
        """,
        (
            payment["student_id"],
            school_id,
            payment["academic_year"],
            payment["term"],
            payment_id,
        ),
    ).fetchone()[0]
    previous_balance = int(charges) - int(other_payments)
    new_balance = previous_balance - int(payment["amount"])
    cursor = conn.execute(
        """
        INSERT INTO receipts
            (school_id, payment_id, receipt_number, previous_balance, new_balance)
        VALUES (?, ?, ?, ?, ?)
        """,
        (
            school_id,
            payment_id,
            f"PENDING-{payment_id}",
            previous_balance,
            new_balance,
        ),
    )
    receipt_id = int(cursor.lastrowid)
    receipt_number = f"RC-{datetime.now().year}-{receipt_id:06d}"
    conn.execute(
        "UPDATE receipts SET receipt_number = ? WHERE id = ?",
        (receipt_number, receipt_id),
    )
    add_audit_log(
        conn,
        school_id,
        actor_id,
        "Receipt generated",
        "receipt",
        receipt_id,
        receipt_number,
    )
    return receipt_id


def embed_school_logo(workbook_stream: BytesIO, logo_url: str | None) -> BytesIO:
    if not logo_url or not logo_url.startswith("/static/"):
        workbook_stream.seek(0)
        return workbook_stream
    logo_path = os.path.join(
        app.root_path, logo_url.lstrip("/").replace("/", os.sep)
    )
    extension = os.path.splitext(logo_path)[1].lower().lstrip(".")
    image_types = {
        "png": "image/png",
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
    }
    if extension not in image_types or not os.path.isfile(logo_path):
        workbook_stream.seek(0)
        return workbook_stream

    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    drawing_ns = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
    drawing_main_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    ET.register_namespace("", main_ns)
    ET.register_namespace("r", rel_ns)
    ET.register_namespace("xdr", drawing_ns)
    ET.register_namespace("a", drawing_main_ns)

    workbook_stream.seek(0)
    output = BytesIO()
    logo_filename = f"school-logo.{extension}"
    with zipfile.ZipFile(workbook_stream, "r") as source:
        worksheet_root = ET.fromstring(source.read("xl/worksheets/sheet1.xml"))
        ET.SubElement(worksheet_root, f"{{{main_ns}}}drawing", {f"{{{rel_ns}}}id": "rId1"})
        worksheet_xml = ET.tostring(
            worksheet_root, encoding="utf-8", xml_declaration=True
        )
        worksheet_relationships = ET.Element(
            f"{{{package_rel_ns}}}Relationships"
        )
        ET.SubElement(
            worksheet_relationships,
            f"{{{package_rel_ns}}}Relationship",
            {
                "Id": "rId1",
                "Type": f"{rel_ns}/drawing",
                "Target": "../drawings/drawing1.xml",
            },
        )

        width_emu = 72 * 9525
        height_emu = 54 * 9525
        anchor = ET.Element(f"{{{drawing_ns}}}wsDr")
        one_cell_anchor = ET.SubElement(anchor, f"{{{drawing_ns}}}oneCellAnchor")
        start = ET.SubElement(one_cell_anchor, f"{{{drawing_ns}}}from")
        for tag, value in (("col", "0"), ("colOff", "0"), ("row", "0"), ("rowOff", "0")):
            ET.SubElement(start, f"{{{drawing_ns}}}{tag}").text = value
        ET.SubElement(
            one_cell_anchor,
            f"{{{drawing_ns}}}ext",
            {"cx": str(width_emu), "cy": str(height_emu)},
        )
        picture = ET.SubElement(one_cell_anchor, f"{{{drawing_ns}}}pic")
        non_visual = ET.SubElement(picture, f"{{{drawing_ns}}}nvPicPr")
        ET.SubElement(
            non_visual,
            f"{{{drawing_ns}}}cNvPr",
            {"id": "1", "name": "School logo"},
        )
        ET.SubElement(non_visual, f"{{{drawing_ns}}}cNvPicPr")
        fill = ET.SubElement(picture, f"{{{drawing_ns}}}blipFill")
        ET.SubElement(
            fill,
            f"{{{drawing_main_ns}}}blip",
            {f"{{{rel_ns}}}embed": "rId1"},
        )
        stretch = ET.SubElement(fill, f"{{{drawing_main_ns}}}stretch")
        ET.SubElement(stretch, f"{{{drawing_main_ns}}}fillRect")
        shape = ET.SubElement(picture, f"{{{drawing_ns}}}spPr")
        transform = ET.SubElement(shape, f"{{{drawing_main_ns}}}xfrm")
        ET.SubElement(
            transform, f"{{{drawing_main_ns}}}off", {"x": "0", "y": "0"}
        )
        ET.SubElement(
            transform,
            f"{{{drawing_main_ns}}}ext",
            {"cx": str(width_emu), "cy": str(height_emu)},
        )
        geometry = ET.SubElement(
            shape,
            f"{{{drawing_main_ns}}}prstGeom",
            {"prst": "rect"},
        )
        ET.SubElement(geometry, f"{{{drawing_main_ns}}}avLst")
        ET.SubElement(one_cell_anchor, f"{{{drawing_ns}}}clientData")
        drawing_xml = ET.tostring(anchor, encoding="utf-8", xml_declaration=True)

        drawing_relationships = ET.Element(
            f"{{{package_rel_ns}}}Relationships"
        )
        ET.SubElement(
            drawing_relationships,
            f"{{{package_rel_ns}}}Relationship",
            {
                "Id": "rId1",
                "Type": f"{rel_ns}/image",
                "Target": f"../media/{logo_filename}",
            },
        )
        content_types = ET.fromstring(source.read("[Content_Types].xml"))
        if not any(
            element.attrib.get("Extension") == extension
            for element in content_types
        ):
            ET.SubElement(
                content_types,
                "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
                {"Extension": extension, "ContentType": image_types[extension]},
            )
        ET.SubElement(
            content_types,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
            {
                "PartName": "/xl/drawings/drawing1.xml",
                "ContentType": "application/vnd.openxmlformats-officedocument.drawing+xml",
            },
        )
        content_types_xml = ET.tostring(
            content_types, encoding="utf-8", xml_declaration=True
        )
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as destination:
            for item in source.infolist():
                if item.filename == "xl/worksheets/sheet1.xml":
                    data = worksheet_xml
                elif item.filename == "[Content_Types].xml":
                    data = content_types_xml
                else:
                    data = source.read(item.filename)
                destination.writestr(item, data)
            destination.writestr(
                "xl/worksheets/_rels/sheet1.xml.rels",
                ET.tostring(
                    worksheet_relationships,
                    encoding="utf-8",
                    xml_declaration=True,
                ),
            )
            destination.writestr(
                "xl/drawings/drawing1.xml", drawing_xml
            )
            destination.writestr(
                "xl/drawings/_rels/drawing1.xml.rels",
                ET.tostring(
                    drawing_relationships,
                    encoding="utf-8",
                    xml_declaration=True,
                ),
            )
            destination.write(logo_path, f"xl/media/{logo_filename}")
    output.seek(0)
    return output


def teacher_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("teacher_logged_in"):
            flash("Please sign in before using the school records system.")
            return redirect(url_for("login"))
        return view(*args, **kwargs)

    return wrapped


def role_required(*allowed_roles: str):
    def decorate(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            if not session.get("teacher_logged_in"):
                flash("Please sign in before using the school records system.")
                return redirect(url_for("login"))
            with get_db() as conn:
                teacher = conn.execute(
                    "SELECT role, is_super_admin FROM teachers WHERE id = ?",
                    (session.get("teacher_id"),),
                ).fetchone()
            if not teacher:
                session.clear()
                flash("Your account is no longer available. Please sign in again.")
                return redirect(url_for("login"))
            session["teacher_role"] = teacher["role"]
            if teacher["is_super_admin"] or teacher["role"] in allowed_roles:
                return view(*args, **kwargs)
            abort(403)

        return wrapped

    return decorate


def teacher_can_manage_school(teacher_id: int | None, school_id: int) -> bool:
    if teacher_id is None:
        return False
    with get_db() as conn:
        teacher = conn.execute("SELECT school_id, is_super_admin FROM teachers WHERE id = ?", (teacher_id,)).fetchone()
    if not teacher:
        return False
    return bool(teacher["is_super_admin"]) or int(teacher["school_id"]) == int(school_id)


def get_active_school_id() -> int:
    try:
        selected_school_id = session.get("school_id")
    except RuntimeError:
        selected_school_id = None

    if selected_school_id is None:
        with get_db() as conn:
            teacher_id = None
            try:
                teacher_id = session.get("teacher_id")
            except RuntimeError:
                teacher_id = None
            teacher = conn.execute("SELECT school_id FROM teachers WHERE id = ?", (teacher_id,)).fetchone() if teacher_id else None
            selected_school_id = teacher["school_id"] if teacher else 1
            try:
                session["school_id"] = selected_school_id
            except RuntimeError:
                pass
    return int(selected_school_id)


def get_current_school() -> Dict[str, Any]:
    school_id = get_active_school_id()
    with get_db() as conn:
        row = conn.execute("SELECT * FROM schools WHERE id = ?", (school_id,)).fetchone()
    return dict(row) if row else {"id": school_id, "name": "School", "motto": "", "logo_url": "/static/sample-report-form.png", "signature_url": "", "report_template": "default"}


def get_report_template_name(grade_label: str) -> str:
    cleaned = str(grade_label).strip().lower().replace(" ", "")
    if cleaned in {"grade6"}:
        return "report_card_grade_6.html"
    if cleaned in {"grade8", "grade9"}:
        return "report_card_grade_8.html" if cleaned == "grade8" else "report_card_grade_9.html"
    return "report_card_default.html"


def get_subject_grade_rubric(grade_level: str, percentage: float) -> Dict[str, Any]:
    bands = [
        (90, "EE1", 8, "Exceeding Expectation 1"),
        (75, "EE2", 7, "Exceeding Expectation 2"),
        (58, "ME1", 6, "Meeting Expectation 1"),
        (41, "ME2", 5, "Meeting Expectation 2"),
        (31, "AE1", 4, "Approaching Expectation 1"),
        (21, "AE2", 3, "Approaching Expectation 2"),
        (11, "BE1", 2, "Below Expectation 1"),
        (0, "BE2", 1, "Below Expectation 2"),
    ]
    for minimum, grade, points, descriptor in bands:
        if percentage >= minimum:
            return {
                "grade": grade,
                "descriptor": descriptor,
                "band": str(points),
                "points": points,
                "band_name": descriptor,
            }
    return {"grade": "BE2", "descriptor": "Below Expectation 2", "band": "1", "points": 1, "band_name": "Below Expectation 2"}


def normalize_grade_level(grade_label: str) -> str:
    cleaned = str(grade_label).strip().lower().replace(" ", "")
    if cleaned in {"playgroup", "pp1", "pp2", "grade1", "grade2", "grade3", "grade4", "grade5"}:
        return "primary"
    if cleaned in {"grade6"}:
        return "kpsea"
    if cleaned in {"grade7", "grade8", "grade9"}:
        return "kjsea"
    return "primary"


def get_grade_info(grade_level: str, percentage: float) -> Dict[str, Any]:
    system = normalize_grade_level(grade_level)
    if system == "kjsea":
        if percentage >= 80:
            return {"grade": "A", "descriptor": "Exceeding Expectation", "band": "4", "band_name": "Exceeding Expectation"}
        if percentage >= 60:
            return {"grade": "B", "descriptor": "Meeting Expectation", "band": "3", "band_name": "Meeting Expectation"}
        if percentage >= 40:
            return {"grade": "C", "descriptor": "Approaching Expectation", "band": "2", "band_name": "Approaching Expectation"}
        return {"grade": "D", "descriptor": "Below Expectation", "band": "1", "band_name": "Below Expectation"}

    if system == "kpsea":
        if percentage >= 80:
            return {"grade": "A", "descriptor": "Exceeding Expectations", "band": "4", "band_name": "Exceeding Expectations"}
        if percentage >= 60:
            return {"grade": "B", "descriptor": "Meeting Expectations", "band": "3", "band_name": "Meeting Expectations"}
        if percentage >= 40:
            return {"grade": "C", "descriptor": "Approaching Expectations", "band": "2", "band_name": "Approaching Expectations"}
        return {"grade": "D", "descriptor": "Below Expectations", "band": "1", "band_name": "Below Expectations"}

    if percentage >= 80:
        return {"grade": "A", "descriptor": "Excellent", "band": "", "band_name": ""}
    if percentage >= 70:
        return {"grade": "B", "descriptor": "Very Good", "band": "", "band_name": ""}
    if percentage >= 60:
        return {"grade": "C", "descriptor": "Good", "band": "", "band_name": ""}
    if percentage >= 50:
        return {"grade": "D", "descriptor": "Satisfactory", "band": "", "band_name": ""}
    return {"grade": "E", "descriptor": "Needs Support", "band": "", "band_name": ""}


def generate_personalized_remark(student_name: str, overall_average: float, strength_subjects: List[str], weak_subjects: List[str]) -> str:
    strengths = ", ".join(strength_subjects[:2]) if strength_subjects else "core learning areas"
    growth = ", ".join(weak_subjects[:2]) if weak_subjects else "key concepts"
    if overall_average >= 80:
        return (
            f"{student_name} is a highly capable and dependable learner who consistently demonstrates excellent understanding. "
            f"The learner shows strength in {strengths} and is encouraged to keep building confidence by extending learning in {growth}."
        )
    if overall_average >= 65:
        return (
            f"{student_name} is a focused and improving learner who is making steady progress. "
            f"The learner performs strongly in {strengths} and can strengthen further by giving extra attention to {growth}."
        )
    if overall_average >= 50:
        return (
            f"{student_name} is a determined learner who shows willingness to grow and improve. "
            f"With continued effort in {growth}, the learner can build on strengths in {strengths} and gain greater confidence."
        )
    return (
        f"{student_name} is a resilient learner who is developing confidence and independence. "
        f"With regular practice and support in {growth}, the learner can make strong progress and continue to grow in {strengths}."
    )


def build_student_report(student_id: int, exam_id: int | None = None) -> Dict[str, Any]:
    school_id = get_active_school_id()
    report_role = session.get("teacher_role", "") if has_request_context() else ""
    report_teacher_id = session.get("teacher_id") if has_request_context() else None
    with get_db() as conn:
        student = conn.execute("SELECT * FROM students WHERE id = ? AND school_id = ?", (student_id, school_id)).fetchone()
        if not student:
            raise ValueError("Student not found")
        marks_rows = conn.execute(
            """
            SELECT m.*, s.name AS subject_name, s.short_name AS short_name,
                   e.name AS exam_name, e.id AS selected_exam_id
            FROM marks m
            JOIN subjects s ON s.id = m.subject_id
            JOIN exams e ON e.id = m.exam_id
            WHERE m.student_id = ?
              AND m.school_id = ?
              AND (? IS NULL OR m.exam_id = ?)
              AND (
                  ? != 'teacher'
                  OR EXISTS (
                      SELECT 1 FROM teacher_learning_areas a
                      WHERE a.school_id = m.school_id
                        AND a.teacher_id = ?
                        AND a.subject_id = m.subject_id
                        AND a.grade = ?
                  )
              )
            ORDER BY s.name
            """,
            (
                student_id,
                school_id,
                exam_id,
                exam_id,
                report_role,
                report_teacher_id,
                student["grade"],
            ),
        ).fetchall()

    subject_scores: Dict[str, List[float]] = {}
    for row in marks_rows:
        out_of = float(row["out_of"] or 100)
        score = float(row["marks_obtained"]) / out_of * 100 if out_of else 0
        subject_scores.setdefault(row["subject_name"], []).append(score)

    subject_summary: List[Dict[str, Any]] = []
    for name, values in sorted(subject_scores.items()):
        average = round(sum(values) / len(values), 1)
        rubric = get_subject_grade_rubric(student["grade"], average)
        subject_summary.append(
            {
                "name": name,
                "average": average,
                "grade": rubric,
                "points": rubric["points"],
            }
        )

    overall_average = round(sum(item["average"] for item in subject_summary) / len(subject_summary), 1) if subject_summary else 0.0
    overall_grade = get_subject_grade_rubric(student["grade"], overall_average)
    mark_entries = []
    for row in marks_rows:
        out_of = float(row["out_of"] or 100)
        percentage = round(float(row["marks_obtained"]) / out_of * 100, 1) if out_of else 0
        rubric = get_subject_grade_rubric(student["grade"], percentage)
        mark_entries.append(
            {
                "name": row["subject_name"],
                "exam_name": row["exam_name"],
                "score": float(row["marks_obtained"]),
                "out_of": out_of,
                "percentage": percentage,
                "points": rubric["points"],
                "grade": rubric["grade"],
            }
        )

    with get_db() as conn:
        class_students = conn.execute(
            """
            SELECT id FROM students
            WHERE grade = ? AND class_name = ? AND school_id = ?
              AND (
                  ? != 'teacher'
                  OR EXISTS (
                      SELECT 1 FROM teacher_learning_areas a
                      WHERE a.school_id = students.school_id
                        AND a.teacher_id = ?
                        AND a.grade = students.grade
                  )
              )
            """,
            (student["grade"], student["class_name"], school_id, report_role, report_teacher_id),
        ).fetchall()
        peer_subject_rows = conn.execute(
            """
            SELECT m.student_id, m.subject_id,
                   AVG(CASE WHEN COALESCE(m.out_of, 100) > 0
                       THEN m.marks_obtained * 100.0 / m.out_of ELSE 0 END) AS average
            FROM marks m
            JOIN students st ON st.id = m.student_id AND st.school_id = m.school_id
            WHERE st.grade = ? AND st.class_name = ? AND st.school_id = ?
              AND (? IS NULL OR m.exam_id = ?)
              AND (
                  ? != 'teacher'
                  OR EXISTS (
                      SELECT 1 FROM teacher_learning_areas a
                      WHERE a.school_id = m.school_id
                        AND a.teacher_id = ?
                        AND a.subject_id = m.subject_id
                        AND a.grade = st.grade
                  )
              )
            GROUP BY m.student_id, m.subject_id
            """,
            (
                student["grade"],
                student["class_name"],
                school_id,
                exam_id,
                exam_id,
                report_role,
                report_teacher_id,
            ),
        ).fetchall()

    averages_by_student: Dict[int, List[float]] = {}
    for row in peer_subject_rows:
        averages_by_student.setdefault(int(row["student_id"]), []).append(float(row["average"]))
    positions = [
        {
            "id": int(roster_student["id"]),
            "average": (
                round(sum(averages_by_student[int(roster_student["id"])]) / len(averages_by_student[int(roster_student["id"])]), 1)
                if averages_by_student.get(int(roster_student["id"]))
                else 0.0
            ),
        }
        for roster_student in class_students
    ]

    positions_sorted = sorted(positions, key=lambda item: item["average"], reverse=True)
    current_position = 1
    for index, item in enumerate(positions_sorted):
        if item["id"] == student_id:
            current_position = index + 1
            break

    strength_subjects = [item["name"] for item in subject_summary if item["average"] >= 75]
    weak_subjects = [item["name"] for item in subject_summary if item["average"] < 41]
    if not weak_subjects:
        weak_subjects = ["core concepts and revision habits"]

    remark = generate_personalized_remark(student["full_name"], overall_average, strength_subjects, weak_subjects)

    return {
        "student": dict(student),
        "subject_summary": subject_summary,
        "mark_entries": mark_entries,
        "overall_average": overall_average,
        "overall_grade": overall_grade,
        "position": current_position,
        "strength_subjects": strength_subjects,
        "weak_subjects": weak_subjects,
        "remark": remark,
        "performance_band": overall_grade["band_name"] or overall_grade["descriptor"],
        "report_template": get_report_template_name(student["grade"]),
        "exam_name": marks_rows[0]["exam_name"] if marks_rows else "All recorded exams",
    }


@app.context_processor
def inject_school_settings():
    teacher = None
    if session.get("teacher_id"):
        with get_db() as conn:
            teacher = conn.execute(
                "SELECT role, is_super_admin FROM teachers WHERE id = ?",
                (session["teacher_id"],),
            ).fetchone()
    return {
        "school_settings": get_current_school(),
        "available_schools": get_available_schools(),
        "current_role": teacher["role"] if teacher else None,
        "is_super_admin": bool(teacher["is_super_admin"]) if teacher else False,
    }


def get_available_schools() -> List[Dict[str, Any]]:
    teacher_id = session.get("teacher_id")
    if not teacher_id:
        with get_db() as conn:
            rows = conn.execute("SELECT id, name, motto, logo_url FROM schools ORDER BY name").fetchall()
        return [dict(row) for row in rows]

    with get_db() as conn:
        teacher = conn.execute("SELECT school_id, is_super_admin FROM teachers WHERE id = ?", (teacher_id,)).fetchone()
    if not teacher:
        return []
    if teacher["is_super_admin"]:
        with get_db() as conn:
            rows = conn.execute("SELECT id, name, motto, logo_url FROM schools ORDER BY name").fetchall()
        return [dict(row) for row in rows]
    with get_db() as conn:
        rows = conn.execute("SELECT id, name, motto, logo_url FROM schools WHERE id = ? ORDER BY name", (teacher["school_id"],)).fetchall()
    return [dict(row) for row in rows]


@app.route("/")
@teacher_required
def index():
    school_id = get_active_school_id()
    with get_db() as conn:
        teacher = conn.execute(
            "SELECT role, is_super_admin FROM teachers WHERE id = ?",
            (session.get("teacher_id"),),
        ).fetchone()
    role = teacher["role"] if teacher else "teacher"
    with get_db() as conn:
        student_count = conn.execute("SELECT COUNT(*) AS count FROM students WHERE school_id = ?", (school_id,)).fetchone()["count"] if role in {"admin", "registrar", "super_admin"} else 0
        subject_count = conn.execute("SELECT COUNT(*) AS count FROM subjects WHERE school_id = ?", (school_id,)).fetchone()["count"] if role in {"admin", "teacher", "exam_officer", "super_admin"} else 0
        exam_count = conn.execute("SELECT COUNT(*) AS count FROM exams WHERE school_id = ?", (school_id,)).fetchone()["count"] if role in {"admin", "exam_officer", "super_admin"} else 0
        mark_count = conn.execute("SELECT COUNT(*) AS count FROM marks WHERE school_id = ?", (school_id,)).fetchone()["count"] if role in {"admin", "teacher", "exam_officer", "super_admin"} else 0
        finance = conn.execute(
            """
            SELECT
                COALESCE((SELECT SUM(amount) FROM payments
                          WHERE school_id = ? AND status = 'matched'
                            AND date(received_at) = date('now', 'localtime')), 0) AS collected_today,
                MAX(
                    COALESCE((SELECT SUM(amount) FROM fee_charges WHERE school_id = ?), 0)
                    - COALESCE((SELECT SUM(amount) FROM payments
                                WHERE school_id = ? AND status = 'matched'), 0),
                    0
                ) AS outstanding,
                (SELECT COUNT(*) FROM payments WHERE school_id = ?) AS payment_count,
                (SELECT COUNT(*) FROM payments WHERE school_id = ? AND status = 'pending') AS unmatched_count,
                (SELECT COUNT(*) FROM receipts WHERE school_id = ?) AS receipt_count
            """,
            (school_id, school_id, school_id, school_id, school_id, school_id),
        ).fetchone()
        teacher_workload = conn.execute(
            """
            SELECT COUNT(*) AS assignments
            FROM teacher_learning_areas WHERE school_id = ? AND teacher_id = ?
            """,
            (school_id, session.get("teacher_id")),
        ).fetchone()["assignments"]
    return render_template(
        "index.html",
        student_count=student_count,
        subject_count=subject_count,
        exam_count=exam_count,
        mark_count=mark_count,
        finance=finance,
        role=role,
        teacher_workload=teacher_workload,
    )


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        with get_db() as conn:
            teacher = conn.execute(
                "SELECT * FROM teachers WHERE username = ? AND password = ?",
                (request.form.get("username"), request.form.get("password")),
            ).fetchone()
        if teacher:
            session["teacher_logged_in"] = True
            session["teacher_id"] = teacher["id"]
            session["teacher_name"] = teacher["full_name"]
            session["teacher_role"] = teacher["role"]
            session["school_id"] = teacher["school_id"]
            flash("Welcome back.")
            return redirect(url_for("index"))
        flash("The teacher credentials were not accepted.")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    flash("You have signed out.")
    return redirect(url_for("login"))


@app.route("/switch-school/<int:school_id>")
@role_required("super_admin")
def switch_school(school_id: int):
    teacher_id = session.get("teacher_id")
    if not teacher_can_manage_school(teacher_id, school_id):
        flash("You are not allowed to switch to that school.")
        return redirect(url_for("index"))
    with get_db() as conn:
        school = conn.execute("SELECT id FROM schools WHERE id = ?", (school_id,)).fetchone()
    if school:
        session["school_id"] = school_id
        flash("School switched successfully.")
    else:
        flash("The selected school could not be found.")
    return redirect(url_for("index"))


@app.route("/students", methods=["GET", "POST"])
@role_required("admin", "registrar")
def students():
    school_id = get_active_school_id()
    if request.method == "POST":
        grade = (request.form.get("grade") or "").strip()
        stage = (request.form.get("school_stage") or stage_for_grade(grade)).strip()
        relationship = (request.form.get("parent_relationship") or "").strip()
        if not valid_student_stage(grade, stage) or relationship not in {"Father", "Mother", "Guardian"}:
            flash("Choose a grade that belongs to the selected school stage and a parent relationship.")
            return redirect(url_for("students"))
        with get_db() as conn:
            cursor = conn.execute(
                """
                INSERT INTO students (admission_number, full_name, gender, grade, school_stage, class_name, dob, parent_name, parent_relationship, parent_contact, school_id, payment_code)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, '')
                """,
                (
                    request.form.get("admission_number"),
                    request.form.get("full_name"),
                    request.form.get("gender"),
                    grade,
                    stage,
                    request.form.get("class_name"),
                    request.form.get("dob"),
                    request.form.get("parent_name"),
                    relationship,
                    request.form.get("parent_contact"),
                    school_id,
                ),
            )
            student_id = int(cursor.lastrowid)
            payment_code = f"S{school_id}-{student_id:06d}"
            conn.execute(
                "UPDATE students SET payment_code = ? WHERE id = ?",
                (payment_code, student_id),
            )
        flash(f"Student added successfully. Payment code: {payment_code}")
        return redirect(url_for("students"))

    with get_db() as conn:
        student_rows = conn.execute("SELECT * FROM students WHERE school_id = ? ORDER BY grade, class_name, full_name", (school_id,)).fetchall()
    return render_template("students.html", students=student_rows)


@app.route("/students/<int:student_id>/edit", methods=["GET", "POST"])
@role_required("admin", "registrar")
def edit_student(student_id: int):
    school_id = get_active_school_id()
    with get_db() as conn:
        student = conn.execute("SELECT * FROM students WHERE id = ? AND school_id = ?", (student_id, school_id)).fetchone()
    if not student:
        flash("Student not found.")
        return redirect(url_for("students"))

    if request.method == "POST":
        grade = (request.form.get("grade") or "").strip()
        stage = (request.form.get("school_stage") or stage_for_grade(grade)).strip()
        relationship = (request.form.get("parent_relationship") or "").strip()
        if not valid_student_stage(grade, stage) or relationship not in {"Father", "Mother", "Guardian"}:
            flash("Choose a grade that belongs to the selected school stage and a parent relationship.")
            return redirect(url_for("edit_student", student_id=student_id))
        with get_db() as conn:
            conn.execute(
                """
                UPDATE students SET admission_number = ?, full_name = ?, gender = ?, grade = ?,
                    school_stage = ?, class_name = ?, dob = ?, parent_name = ?,
                    parent_relationship = ?, parent_contact = ?
                WHERE id = ? AND school_id = ?
                """,
                (
                    request.form.get("admission_number"),
                    request.form.get("full_name"),
                    request.form.get("gender"),
                    grade,
                    stage,
                    request.form.get("class_name"),
                    request.form.get("dob"),
                    request.form.get("parent_name"),
                    relationship,
                    request.form.get("parent_contact"),
                    student_id,
                    school_id,
                ),
            )
        flash("Student updated successfully.")
        return redirect(url_for("students"))

    return render_template("edit_student.html", student=student)


@app.route("/fees", methods=["GET", "POST"])
@role_required("admin", "finance_officer")
def fees():
    school_id = get_active_school_id()
    if request.method == "POST":
        grade = (request.form.get("grade") or "").strip()
        academic_year = (request.form.get("academic_year") or "").strip()
        term = (request.form.get("term") or "").strip()
        name = (request.form.get("name") or "").strip()
        amount = parse_ksh_amount(request.form.get("amount"))
        if (
            not all((grade, academic_year, name))
            or term not in {"Term 1", "Term 2", "Term 3"}
            or amount is None
        ):
            flash("Enter a grade, year, valid term, fee name, and a positive whole-KSh amount.")
            return redirect(url_for("fees"))

        with get_db() as conn:
            exists = conn.execute(
                """
                SELECT id FROM fee_items
                WHERE school_id = ? AND grade = ? AND academic_year = ? AND term = ? AND name = ?
                """,
                (school_id, grade, academic_year, term, name),
            ).fetchone()
            if exists:
                flash("That fee item already exists for this grade, year, and term.")
                return redirect(url_for("fees"))

            cursor = conn.execute(
                """
                INSERT INTO fee_items (school_id, grade, academic_year, term, name, amount)
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (school_id, grade, academic_year, term, name, amount),
            )
            fee_item_id = int(cursor.lastrowid)
            roster = conn.execute(
                "SELECT id FROM students WHERE school_id = ? AND grade = ?",
                (school_id, grade),
            ).fetchall()
            for student in roster:
                conn.execute(
                    """
                    INSERT OR IGNORE INTO fee_charges
                        (school_id, student_id, fee_item_id, item_name, academic_year, term, amount)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    (school_id, student["id"], fee_item_id, name, academic_year, term, amount),
                )
            add_audit_log(
                conn,
                school_id,
                session.get("teacher_id"),
                "Fee item created and charged",
                "fee_item",
                fee_item_id,
                f"{grade} {academic_year} {term}: {name}, KSh {amount}",
            )
        flash(f"Fee item saved and charged to {len(roster)} {grade} learner(s).")
        return redirect(url_for("fees"))

    with get_db() as conn:
        fee_items = conn.execute(
            """
            SELECT f.*, COUNT(c.id) AS learners_charged
            FROM fee_items f LEFT JOIN fee_charges c ON c.fee_item_id = f.id
            WHERE f.school_id = ?
            GROUP BY f.id
            ORDER BY f.academic_year DESC, f.term, f.grade, f.name
            """,
            (school_id,),
        ).fetchall()
        grades = conn.execute(
            "SELECT DISTINCT grade FROM students WHERE school_id = ? ORDER BY grade",
            (school_id,),
        ).fetchall()
    return render_template("fees.html", fee_items=fee_items, grades=grades)


@app.route("/fees/<int:fee_item_id>/apply", methods=["POST"])
@role_required("admin", "finance_officer")
def apply_fee_item(fee_item_id: int):
    school_id = get_active_school_id()
    with get_db() as conn:
        fee_item = conn.execute(
            "SELECT * FROM fee_items WHERE id = ? AND school_id = ?",
            (fee_item_id, school_id),
        ).fetchone()
        if not fee_item:
            flash("Fee item not found.")
            return redirect(url_for("fees"))
        cursor = conn.execute(
            """
            INSERT OR IGNORE INTO fee_charges
                (school_id, student_id, fee_item_id, item_name, academic_year, term, amount)
            SELECT ?, id, ?, ?, ?, ?, ?
            FROM students WHERE school_id = ? AND grade = ?
            """,
            (
                school_id,
                fee_item_id,
                fee_item["name"],
                fee_item["academic_year"],
                fee_item["term"],
                fee_item["amount"],
                school_id,
                fee_item["grade"],
            ),
        )
        add_audit_log(
            conn,
            school_id,
            session.get("teacher_id"),
            "Fee item applied to uncharged learners",
            "fee_item",
            fee_item_id,
            f"Added {cursor.rowcount} learner charge(s).",
        )
    flash(f"Fee item applied to {cursor.rowcount} newly eligible learner(s).")
    return redirect(url_for("fees"))


@app.route("/payments", methods=["GET", "POST"])
@role_required("admin", "finance_officer")
def payments():
    school_id = get_active_school_id()
    if request.method == "POST":
        amount = parse_ksh_amount(request.form.get("amount"))
        channel = (request.form.get("channel") or "").strip()
        student_id = request.form.get("student_id", type=int)
        academic_year = (request.form.get("academic_year") or "").strip()
        term = (request.form.get("term") or "").strip()
        transaction_code = (request.form.get("transaction_code") or "").strip()
        payer_phone = (request.form.get("payer_phone") or "").strip()
        if (
            amount is None
            or channel not in {"M-Pesa", "Bank", "Cash", "Other"}
            or not student_id
            or not academic_year
            or not term
            or not re.fullmatch(r"0\d{9}", payer_phone)
        ):
            flash("Choose a learner, fee year and term, and enter a valid 10-digit Kenyan phone number and positive amount.")
            return redirect(url_for("payments"))

        with get_db() as conn:
            candidate = conn.execute(
                """
                SELECT * FROM students WHERE id = ? AND school_id = ?
                """,
                (student_id, school_id),
            ).fetchone()
            if not candidate:
                flash("The selected learner could not be found.")
                return redirect(url_for("payments"))
            applicable_charges = conn.execute(
                """
                SELECT COUNT(*) FROM fee_charges
                WHERE school_id = ? AND student_id = ? AND academic_year = ? AND term = ?
                """,
                (school_id, student_id, academic_year, term),
            ).fetchone()[0]
            if not applicable_charges:
                flash("The selected learner has no fee charges for that year and term.")
                return redirect(url_for("payments"))
            if transaction_code and conn.execute(
                "SELECT id FROM payments WHERE school_id = ? AND transaction_code = ?",
                (school_id, transaction_code),
            ).fetchone():
                flash("That transaction code has already been recorded.")
                return redirect(url_for("payments"))

            phone_matches = bool(
                candidate["parent_contact"]
                and normalize_phone(payer_phone) == normalize_phone(candidate["parent_contact"])
            )
            auto_match = phone_matches
            status = "matched" if auto_match else "pending"
            timestamp = datetime.now().isoformat(sep=" ", timespec="seconds")
            cursor = conn.execute(
                """
                INSERT INTO payments
                    (school_id, amount, channel, transaction_code, payer_phone, reference,
                     academic_year, term, student_id, suggested_student_id, status, received_at, verified_by,
                     verified_at, verification_note)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    school_id,
                    amount,
                    channel,
                    transaction_code or None,
                    payer_phone or None,
                    candidate["payment_code"] or candidate["admission_number"],
                    academic_year,
                    term,
                    candidate["id"] if auto_match else None,
                    candidate["id"] if candidate else None,
                    status,
                    timestamp,
                    session.get("teacher_id") if auto_match else None,
                    timestamp if auto_match else None,
                    "Matched using learner reference and payer details." if auto_match else None,
                ),
            )
            payment_id = int(cursor.lastrowid)
            if auto_match:
                receipt_id = create_payment_receipt(
                    conn, payment_id, school_id, session.get("teacher_id")
                )
                add_audit_log(
                    conn,
                    school_id,
                    session.get("teacher_id"),
                    "Payment matched",
                    "payment",
                    payment_id,
                    f"KSh {amount} assigned to {candidate['full_name']}",
                )
            else:
                receipt_id = None
                add_audit_log(
                    conn,
                    school_id,
                    session.get("teacher_id"),
                    "Payment requires verification",
                    "payment",
                    payment_id,
                    f"KSh {amount}; learner {candidate['full_name']}; {academic_year} {term}",
                )
        if receipt_id:
            flash("Payment matched and receipt generated.")
            return redirect(url_for("payment_receipt", receipt_id=receipt_id))
        flash("Payment recorded in the unmatched queue for bursar verification.")
        return redirect(url_for("payments", queue="unmatched"))

    with get_db() as conn:
        payment_rows = conn.execute(
            """
            SELECT p.*, s.full_name AS student_name, s.admission_number,
                   suggested.full_name AS suggested_name,
                   r.id AS receipt_id, r.receipt_number
            FROM payments p
            LEFT JOIN students s ON s.id = p.student_id
            LEFT JOIN students suggested ON suggested.id = p.suggested_student_id
            LEFT JOIN receipts r ON r.payment_id = p.id
            WHERE p.school_id = ?
            ORDER BY p.received_at DESC, p.id DESC
            """,
            (school_id,),
        ).fetchall()
        student_rows = conn.execute(
            "SELECT id, full_name, admission_number FROM students WHERE school_id = ? ORDER BY full_name",
            (school_id,),
        ).fetchall()
        fee_periods = conn.execute(
            """
            SELECT DISTINCT academic_year, term
            FROM fee_charges WHERE school_id = ?
            ORDER BY academic_year DESC, term
            """,
            (school_id,),
        ).fetchall()
    return render_template(
        "payments.html",
        payments=payment_rows,
        students=student_rows,
        fee_periods=fee_periods,
        unmatched_only=request.args.get("queue") == "unmatched",
    )


@app.route("/payments/<int:payment_id>/verify", methods=["POST"])
@role_required("admin", "finance_officer")
def verify_payment(payment_id: int):
    school_id = get_active_school_id()
    student_id = request.form.get("student_id", type=int)
    note = (request.form.get("verification_note") or "").strip()
    if not student_id or not note:
        flash("Select the learner and record a verification note.")
        return redirect(url_for("payments", queue="unmatched"))

    with get_db() as conn:
        payment = conn.execute(
            "SELECT * FROM payments WHERE id = ? AND school_id = ? AND status = 'pending'",
            (payment_id, school_id),
        ).fetchone()
        student = conn.execute(
            "SELECT id FROM students WHERE id = ? AND school_id = ?",
            (student_id, school_id),
        ).fetchone()
        if not payment or not student:
            flash("The pending payment or selected learner could not be found.")
            return redirect(url_for("payments", queue="unmatched"))
        timestamp = datetime.now().isoformat(sep=" ", timespec="seconds")
        conn.execute(
            """
            UPDATE payments
            SET student_id = ?, status = 'matched', verified_by = ?, verified_at = ?,
                verification_note = ?
            WHERE id = ?
            """,
            (student_id, session.get("teacher_id"), timestamp, note, payment_id),
        )
        receipt_id = create_payment_receipt(
            conn, payment_id, school_id, session.get("teacher_id")
        )
        add_audit_log(
            conn,
            school_id,
            session.get("teacher_id"),
            "Payment verified",
            "payment",
            payment_id,
            note,
        )
    flash("Payment verified and receipt generated.")
    return redirect(url_for("payment_receipt", receipt_id=receipt_id))


@app.route("/receipts/<int:receipt_id>")
@role_required("admin", "finance_officer")
def payment_receipt(receipt_id: int):
    school_id = get_active_school_id()
    with get_db() as conn:
        receipt = conn.execute(
            """
            SELECT r.*, p.amount, p.channel, p.transaction_code, p.payer_phone,
                   p.received_at, p.reference, p.academic_year, p.term,
                   r.previous_balance, r.new_balance,
                   s.full_name, s.admission_number,
                   s.id AS student_id, s.grade, s.payment_code, school.name AS school_name
            FROM receipts r
            JOIN payments p ON p.id = r.payment_id
            JOIN students s ON s.id = p.student_id
            JOIN schools school ON school.id = r.school_id
            WHERE r.id = ? AND r.school_id = ? AND p.status = 'matched'
            """,
            (receipt_id, school_id),
        ).fetchone()
        if not receipt:
            flash("Receipt not found.")
            return redirect(url_for("payments"))
    return render_template("payment_receipt.html", receipt=receipt)


@app.route("/students/<int:student_id>/statement")
@role_required("admin", "finance_officer")
def fee_statement(student_id: int):
    school_id = get_active_school_id()
    with get_db() as conn:
        student = conn.execute(
            "SELECT * FROM students WHERE id = ? AND school_id = ?",
            (student_id, school_id),
        ).fetchone()
        if not student:
            flash("Student not found.")
            return redirect(url_for("index"))
        charges = conn.execute(
            """
            SELECT id, item_name AS description, amount, created_at, academic_year, term
            FROM fee_charges WHERE student_id = ? AND school_id = ?
            """,
            (student_id, school_id),
        ).fetchall()
        credits = conn.execute(
            """
            SELECT p.id, p.amount, p.received_at AS created_at, p.channel,
                   p.transaction_code, p.academic_year, p.term,
                   r.id AS receipt_id, r.receipt_number
            FROM payments p LEFT JOIN receipts r ON r.payment_id = p.id
            WHERE p.student_id = ? AND p.school_id = ? AND p.status = 'matched'
            """,
            (student_id, school_id),
        ).fetchall()

    entries: List[Dict[str, Any]] = []
    for charge in charges:
        entries.append(
            {
                "id": charge["id"],
                "created_at": charge["created_at"],
                "description": f"{charge['description']} ({charge['academic_year']} {charge['term']})",
                "debit": int(charge["amount"]),
                "credit": 0,
                "receipt_id": None,
            }
        )
    for credit in credits:
        entries.append(
            {
                "id": credit["id"],
                "created_at": credit["created_at"],
                "description": f"{credit['channel']} payment"
                + (
                    f" ({credit['academic_year']} {credit['term']})"
                    if credit["academic_year"] and credit["term"]
                    else ""
                )
                + (f" · {credit['transaction_code']}" if credit["transaction_code"] else ""),
                "debit": 0,
                "credit": int(credit["amount"]),
                "receipt_id": credit["receipt_id"],
            }
        )
    entries.sort(key=lambda entry: (entry["created_at"], entry["id"]))
    balance = 0
    for entry in entries:
        balance += entry["debit"] - entry["credit"]
        entry["balance"] = balance
    total_charged = sum(entry["debit"] for entry in entries)
    total_paid = sum(entry["credit"] for entry in entries)
    return render_template(
        "fee_statement.html",
        student=student,
        entries=entries,
        total_charged=total_charged,
        total_paid=total_paid,
        balance=balance,
    )


@app.route("/teachers", methods=["GET", "POST"])
@role_required("admin")
def teachers():
    school_id = get_active_school_id()
    if request.method == "POST":
        full_name = (request.form.get("full_name") or "").strip()
        username = (request.form.get("username") or "").strip()
        password = request.form.get("password") or ""
        role = (request.form.get("role") or "").strip()
        if not full_name or not username or not password or role not in {
            "finance_officer",
            "registrar",
            "exam_officer",
            "teacher",
        }:
            flash("Enter the teacher's name, username, password, and a valid school role.")
            return redirect(url_for("teachers"))
        with get_db() as conn:
            conn.execute(
                "INSERT INTO teachers (username, password, full_name, school_id, role) VALUES (?, ?, ?, ?, ?)",
                (username, password, full_name, school_id, role),
            )
        flash("Teacher account added successfully.")
        return redirect(url_for("teachers"))

    with get_db() as conn:
        teacher_rows = conn.execute(
            """
            SELECT t.*, COUNT(a.id) AS workload
            FROM teachers t
            LEFT JOIN teacher_learning_areas a ON a.teacher_id = t.id AND a.school_id = t.school_id
            WHERE t.school_id = ?
            GROUP BY t.id
            ORDER BY t.full_name
            """,
            (school_id,),
        ).fetchall()
    return render_template("teachers.html", teachers=teacher_rows)


@app.route("/subjects", methods=["GET", "POST"])
@role_required("admin")
def subjects():
    school_id = get_active_school_id()
    if request.method == "POST":
        name = (request.form.get("name") or "").strip()
        short_name = (request.form.get("short_name") or "").strip()
        grade_level = (request.form.get("grade_level") or "").strip()
        if not name or not grade_level:
            flash("Enter a learning area name and grade.")
            return redirect(url_for("subjects"))
        with get_db() as conn:
            conn.execute(
                "INSERT INTO subjects (name, short_name, grade_level, school_id) VALUES (?, ?, ?, ?)",
                (name, short_name, grade_level, school_id),
            )
        flash("Learning area added successfully.")
        return redirect(url_for("subjects"))

    with get_db() as conn:
        subject_rows = conn.execute("SELECT * FROM subjects WHERE school_id = ? ORDER BY name", (school_id,)).fetchall()
        teacher_rows = conn.execute(
            "SELECT id, full_name FROM teachers WHERE school_id = ? ORDER BY full_name",
            (school_id,),
        ).fetchall()
        assignment_rows = conn.execute(
            """
            SELECT a.id, a.grade, t.full_name, s.name
            FROM teacher_learning_areas a
            JOIN teachers t ON t.id = a.teacher_id
            JOIN subjects s ON s.id = a.subject_id
            WHERE a.school_id = ?
            ORDER BY t.full_name, a.grade, s.name
            """,
            (school_id,),
        ).fetchall()
    return render_template(
        "subjects.html",
        subjects=subject_rows,
        teachers=teacher_rows,
        assignments=assignment_rows,
    )


@app.route("/teacher-assignments", methods=["POST"])
@role_required("admin")
def teacher_assignments():
    school_id = get_active_school_id()
    teacher_id = request.form.get("teacher_id", type=int)
    subject_id = request.form.get("subject_id", type=int)
    grade = (request.form.get("grade") or "").strip()
    if not teacher_id or not subject_id or not grade:
        flash("Choose a teacher, learning area, and grade.")
        return redirect(url_for("subjects"))
    with get_db() as conn:
        teacher = conn.execute(
            "SELECT id FROM teachers WHERE id = ? AND school_id = ?",
            (teacher_id, school_id),
        ).fetchone()
        subject = conn.execute(
            "SELECT id FROM subjects WHERE id = ? AND school_id = ?",
            (subject_id, school_id),
        ).fetchone()
        if not teacher or not subject:
            flash("The selected teacher or learning area does not belong to this school.")
            return redirect(url_for("subjects"))
        conn.execute(
            """
            INSERT OR IGNORE INTO teacher_learning_areas
                (school_id, teacher_id, subject_id, grade)
            VALUES (?, ?, ?, ?)
            """,
            (school_id, teacher_id, subject_id, grade),
        )
    flash("Teacher learning-area assignment saved.")
    return redirect(url_for("subjects"))


@app.route("/exams", methods=["GET", "POST"])
@role_required("admin", "exam_officer")
def exams():
    school_id = get_active_school_id()
    if request.method == "POST":
        grade_level = (request.form.get("grade_level") or "").strip()
        exam_name = (request.form.get("name") or "").strip()
        exam_type = (request.form.get("exam_type") or "").strip()
        term = (request.form.get("term") or "").strip()
        if not exam_name or not exam_type or not term or grade_level not in {
            "All",
            "Playgroup",
            "PP1",
            "PP2",
            *(f"Grade {grade}" for grade in range(1, 13)),
        }:
            flash("Enter an exam name, type, term, and valid grade.")
            return redirect(url_for("exams"))
        with get_db() as conn:
            conn.execute(
                "INSERT INTO exams (name, exam_type, term, grade_level, school_id) VALUES (?, ?, ?, ?, ?)",
                (
                    exam_name,
                    exam_type,
                    term,
                    grade_level,
                    school_id,
                ),
            )
        flash("Exam created successfully.")
        return redirect(url_for("exams"))

    with get_db() as conn:
        exam_rows = conn.execute("SELECT * FROM exams WHERE school_id = ? ORDER BY created_at DESC", (school_id,)).fetchall()
    return render_template("exams.html", exams=exam_rows)


def make_exam_marks_workbook(exam_id: int, school_id: int) -> BytesIO:
    with get_db() as conn:
        exam = conn.execute(
            "SELECT * FROM exams WHERE id = ? AND school_id = ?",
            (exam_id, school_id),
        ).fetchone()
        if not exam:
            raise ValueError("Exam not found.")
        students = conn.execute(
            """
            SELECT * FROM students
            WHERE school_id = ? AND (? = 'All' OR grade = ?)
            ORDER BY grade, class_name, full_name
            """,
            (school_id, exam["grade_level"], exam["grade_level"]),
        ).fetchall()
        learning_areas = conn.execute(
            """
            SELECT * FROM subjects
            WHERE school_id = ? AND (grade_level = 'All' OR grade_level = ?)
            ORDER BY name
            """,
            (school_id, exam["grade_level"]),
        ).fetchall()
    school = get_current_school()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Marks"
    sheet["A1"] = school["name"]
    sheet["B1"] = f"{school['name']} — Assessment Marks"
    sheet["B2"] = f"{exam['name']} · {exam['exam_type']}"
    sheet["B3"] = f"{exam['term']} · {exam['grade_level']}"
    for cell in ("B1", "B2", "B3"):
        sheet[cell].font = Font(bold=True, size=14 if cell == "B1" else 11)
    headers = ["Learner Name", "Admission Number", "Learning Area", "Score (0-100)", "CBE Grade", "Points"]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=6, column=column, value=header)
        cell.font = Font(bold=True)
    current_row = 7
    for student in students:
        for learning_area in learning_areas:
            if learning_area["grade_level"] not in {"All", student["grade"]}:
                continue
            sheet.cell(current_row, 1, student["full_name"])
            sheet.cell(current_row, 2, student["admission_number"])
            sheet.cell(current_row, 3, learning_area["name"])
            sheet.cell(current_row, 4)
            score_ref = f"D{current_row}"
            sheet.cell(
                current_row,
                5,
                f'=IF({score_ref}>=90,"EE1",IF({score_ref}>=75,"EE2",IF({score_ref}>=58,"ME1",IF({score_ref}>=41,"ME2",IF({score_ref}>=31,"AE1",IF({score_ref}>=21,"AE2",IF({score_ref}>=11,"BE1","BE2")))))))',
            )
            sheet.cell(
                current_row,
                6,
                f'=IF({score_ref}>=90,8,IF({score_ref}>=75,7,IF({score_ref}>=58,6,IF({score_ref}>=41,5,IF({score_ref}>=31,4,IF({score_ref}>=21,3,IF({score_ref}>=11,2,1)))))))',
            )
            current_row += 1
    for column, width in {"A": 28, "B": 20, "C": 30, "D": 16, "E": 18, "F": 12}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A7"
    sheet.auto_filter.ref = f"A6:F{max(6, current_row - 1)}"
    workbook_stream = BytesIO()
    workbook.save(workbook_stream)
    return embed_school_logo(workbook_stream, school.get("logo_url"))


@app.route("/exams/<int:exam_id>/marks-template")
@role_required("admin", "exam_officer")
def download_marks_template(exam_id: int):
    school_id = get_active_school_id()
    try:
        workbook = make_exam_marks_workbook(exam_id, school_id)
    except ValueError:
        flash("Exam not found.")
        return redirect(url_for("exams"))
    return send_file(
        workbook,
        as_attachment=True,
        download_name=f"marks-template-{exam_id}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/marks", methods=["GET", "POST"])
@role_required("admin", "exam_officer", "teacher")
def marks():
    school_id = get_active_school_id()
    teacher_id = session.get("teacher_id")
    role = session.get("teacher_role")
    assigned_subjects = None
    with get_db() as conn:
        if role == "teacher":
            assignments = conn.execute(
                """
                SELECT grade, subject_id FROM teacher_learning_areas
                WHERE school_id = ? AND teacher_id = ?
                """,
                (school_id, teacher_id),
            ).fetchall()
            assigned_subjects = {
                (assignment["grade"], assignment["subject_id"])
                for assignment in assignments
            }
            students = conn.execute(
                """
                SELECT DISTINCT st.* FROM students st
                JOIN teacher_learning_areas a
                  ON a.school_id = st.school_id AND a.grade = st.grade
                WHERE st.school_id = ? AND a.teacher_id = ?
                ORDER BY st.grade, st.class_name, st.full_name
                """,
                (school_id, teacher_id),
            ).fetchall()
            subjects = conn.execute(
                """
                SELECT DISTINCT s.* FROM subjects s
                JOIN teacher_learning_areas a ON a.subject_id = s.id
                WHERE a.school_id = ? AND a.teacher_id = ?
                ORDER BY s.name
                """,
                (school_id, teacher_id),
            ).fetchall()
            exams = conn.execute(
                """
                SELECT DISTINCT e.* FROM exams e
                WHERE e.school_id = ? AND (e.grade_level = 'All' OR EXISTS (
                    SELECT 1 FROM teacher_learning_areas a
                    WHERE a.school_id = e.school_id AND a.teacher_id = ?
                      AND a.grade = e.grade_level
                ))
                ORDER BY e.created_at DESC
                """,
                (school_id, teacher_id),
            ).fetchall()
        else:
            students = conn.execute("SELECT * FROM students WHERE school_id = ? ORDER BY grade, class_name, full_name", (school_id,)).fetchall()
            subjects = conn.execute("SELECT * FROM subjects WHERE school_id = ? ORDER BY name", (school_id,)).fetchall()
            exams = conn.execute("SELECT * FROM exams WHERE school_id = ? ORDER BY created_at DESC", (school_id,)).fetchall()
            assigned_subjects = {
                (student["grade"], subject["id"])
                for student in students
                for subject in subjects
            }

    if request.method == "POST":
        if "marks_file" in request.files and request.files["marks_file"].filename:
            file = request.files["marks_file"]
            if not file.filename.lower().endswith(".xlsx"):
                flash("Marks import accepts an .xlsx workbook. Download the exam template or export a Google Sheet as .xlsx.")
                return redirect(url_for("marks"))
            exam_id = request.form.get("exam_id", type=int)
            exam = next((item for item in exams if item["id"] == exam_id), None)
            if not exam:
                flash("Select a valid exam for this school.")
                return redirect(url_for("marks"))
            wb = load_workbook(file.stream, data_only=True, read_only=True)
            sheet = wb.active
            rows = list(sheet.iter_rows(values_only=True))
            header_index = next(
                (
                    index
                    for index, row in enumerate(rows)
                    if row and any(str(cell or "").strip().lower() in {"score", "score (0-100)", "marks", "marks obtained"} for cell in row)
                ),
                None,
            )
            if header_index is None:
                flash("The workbook needs a header row with learner, learning area, and score columns.")
                return redirect(url_for("marks"))
            header = [str(cell or "").strip().lower() for cell in rows[header_index]]
            score_index = next((i for i, name in enumerate(header) if name in {"score", "marks", "marks obtained", "score (0-100)"}), None)
            area_index = next((i for i, name in enumerate(header) if name in {"learning area", "subject"}), None)
            admission_index = next((i for i, name in enumerate(header) if name in {"admission", "admission number", "admission_number"}), None)
            name_index = next((i for i, name in enumerate(header) if name in {"learner name", "student name", "name"}), None)
            if score_index is None or area_index is None or (admission_index is None and name_index is None):
                flash("The workbook needs learner name or admission, learning area, and score columns.")
                return redirect(url_for("marks"))

            by_admission = {str(student["admission_number"]).strip().casefold(): student for student in students}
            by_name: Dict[str, List[Any]] = {}
            for student in students:
                by_name.setdefault(str(student["full_name"]).strip().casefold(), []).append(student)
            subject_by_name = {str(subject["name"]).strip().casefold(): subject for subject in subjects}
            imported = 0
            skipped = 0
            invalid_rows = []
            with get_db() as conn:
                for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
                    if not row or not any(value not in (None, "") for value in row):
                        continue
                    try:
                        raw_score = row[score_index] if score_index < len(row) else None
                        score = float(raw_score)
                        if score < 0 or score > 100:
                            raise ValueError
                    except (TypeError, ValueError):
                        skipped += 1
                        invalid_rows.append(str(row_number))
                        continue
                    admission = (
                        str(row[admission_index]).strip().casefold()
                        if admission_index is not None and admission_index < len(row) and row[admission_index]
                        else ""
                    )
                    learner_name = (
                        str(row[name_index]).strip().casefold()
                        if name_index is not None and name_index < len(row) and row[name_index]
                        else ""
                    )
                    area_name = (
                        str(row[area_index]).strip().casefold()
                        if area_index < len(row) and row[area_index]
                        else ""
                    )
                    student = by_admission.get(admission) if admission else None
                    if not student and learner_name:
                        name_matches = by_name.get(learner_name, [])
                        student = name_matches[0] if len(name_matches) == 1 else None
                    subject = subject_by_name.get(area_name)
                    if (
                        not student
                        or not subject
                        or (learner_name and learner_name != student["full_name"].strip().casefold())
                        or (exam["grade_level"] != "All" and student["grade"] != exam["grade_level"])
                        or subject["grade_level"] not in {"All", student["grade"]}
                    ):
                        skipped += 1
                        invalid_rows.append(str(row_number))
                        continue
                    if role == "teacher":
                        if (student["grade"], subject["id"]) not in assigned_subjects:
                            skipped += 1
                            invalid_rows.append(str(row_number))
                            continue
                    updated = conn.execute(
                        """
                        UPDATE marks SET marks_obtained = ?, out_of = 100
                        WHERE student_id = ? AND exam_id = ? AND subject_id = ? AND school_id = ?
                        """,
                        (score, student["id"], exam_id, subject["id"], school_id),
                    )
                    if not updated.rowcount:
                        conn.execute(
                            """
                            INSERT INTO marks (student_id, exam_id, subject_id, marks_obtained, out_of, school_id)
                            VALUES (?, ?, ?, ?, 100, ?)
                            """,
                            (student["id"], exam_id, subject["id"], score, school_id),
                        )
                    imported += 1
            if skipped:
                flash(f"Imported {imported} mark(s); skipped {skipped} invalid or unassigned row(s): {', '.join(invalid_rows[:10])}.")
            else:
                flash(f"Imported {imported} mark(s) from the workbook.")
            wb.close()
            return redirect(url_for("marks"))

        exam_id = request.form.get("exam_id", type=int)
        if not exam_id or not any(item["id"] == exam_id for item in exams):
            flash("Select a valid exam before saving marks.")
            return redirect(url_for("marks"))
        exam = next(item for item in exams if item["id"] == exam_id)
        with get_db() as conn:
            for student in students:
                for subject in subjects:
                    raw_value = request.form.get(f"mark_{student['id']}_{subject['id']}")
                    if raw_value is None or raw_value == "":
                        continue
                    if (student["grade"], subject["id"]) not in assigned_subjects:
                        continue
                    try:
                        score = float(raw_value)
                    except ValueError:
                        conn.rollback()
                        flash(f"Invalid score for {student['full_name']} in {subject['name']}.")
                        return redirect(url_for("marks"))
                    if score < 0 or score > 100:
                        conn.rollback()
                        flash(f"Score for {student['full_name']} in {subject['name']} must be 0 to 100.")
                        return redirect(url_for("marks"))
                    if (exam["grade_level"] != "All" and student["grade"] != exam["grade_level"]) or subject["grade_level"] not in {"All", student["grade"]}:
                        continue
                    updated = conn.execute(
                        """
                        UPDATE marks SET marks_obtained = ?, out_of = 100
                        WHERE student_id = ? AND exam_id = ? AND subject_id = ? AND school_id = ?
                        """,
                        (score, student["id"], exam_id, subject["id"], school_id),
                    )
                    if updated.rowcount:
                        continue
                    conn.execute(
                        """
                        INSERT INTO marks (student_id, exam_id, subject_id, marks_obtained, out_of, school_id)
                        VALUES (?, ?, ?, ?, 100, ?)
                        """,
                        (student["id"], exam_id, subject["id"], score, school_id),
                    )
        flash("Marks were saved.")
        return redirect(url_for("marks"))

    return render_template(
        "marks.html",
        students=students,
        subjects=subjects,
        exams=exams,
        assigned_subjects=assigned_subjects,
    )


@app.route("/reports")
@role_required("admin", "exam_officer", "teacher")
def reports():
    school_id = get_active_school_id()
    exam_id = request.args.get("exam_id", type=int)
    with get_db() as conn:
        exams = conn.execute(
            "SELECT id, name, exam_type, term, grade_level FROM exams WHERE school_id = ? ORDER BY created_at DESC",
            (school_id,),
        ).fetchall()
        if exam_id and not any(exam["id"] == exam_id for exam in exams):
            flash("The selected exam could not be found.")
            return redirect(url_for("reports"))
        if session.get("teacher_role") == "teacher":
            student_rows = conn.execute(
                """
                SELECT DISTINCT st.* FROM students st
                JOIN teacher_learning_areas a
                  ON a.school_id = st.school_id AND a.grade = st.grade
                WHERE st.school_id = ? AND a.teacher_id = ?
                ORDER BY st.grade, st.class_name, st.full_name
                """,
                (school_id, session.get("teacher_id")),
            ).fetchall()
        else:
            student_rows = conn.execute(
                "SELECT * FROM students WHERE school_id = ? ORDER BY grade, class_name, full_name",
                (school_id,),
            ).fetchall()
    student_reports = [
        build_student_report(student["id"], exam_id) for student in student_rows
    ]
    return render_template(
        "reports.html",
        student_reports=student_reports,
        exams=exams,
        selected_exam_id=exam_id,
    )


def ensure_teacher_student_access(student_id: int) -> None:
    with get_db() as conn:
        student = conn.execute(
            "SELECT grade FROM students WHERE id = ? AND school_id = ?",
            (student_id, get_active_school_id()),
        ).fetchone()
        if not student:
            abort(404)
        if session.get("teacher_role") != "teacher":
            return
        assigned = conn.execute(
            """
            SELECT 1 FROM students st
            JOIN teacher_learning_areas a
              ON a.school_id = st.school_id AND a.grade = st.grade
            WHERE st.id = ? AND st.school_id = ? AND a.teacher_id = ?
            LIMIT 1
            """,
            (student_id, get_active_school_id(), session.get("teacher_id")),
        ).fetchone()
    if not assigned:
        abort(404)


def validate_report_exam(exam_id: int | None) -> None:
    if exam_id is None:
        return
    with get_db() as conn:
        exam = conn.execute(
            "SELECT 1 FROM exams WHERE id = ? AND school_id = ?",
            (exam_id, get_active_school_id()),
        ).fetchone()
    if not exam:
        abort(404)


@app.route("/report/<int:student_id>")
@role_required("admin", "exam_officer", "teacher")
def report_card(student_id: int):
    ensure_teacher_student_access(student_id)
    exam_id = request.args.get("exam_id", type=int)
    report_view = request.args.get("view", "marks")
    if report_view not in {"marks", "points"}:
        abort(400)
    validate_report_exam(exam_id)
    report = build_student_report(student_id, exam_id)
    return render_template(
        "report_card_cbe.html",
        report=report,
        report_view=report_view,
        selected_exam_id=exam_id,
    )


@app.route("/print-report/<int:student_id>")
@role_required("admin", "exam_officer", "teacher")
def print_report(student_id: int):
    ensure_teacher_student_access(student_id)
    exam_id = request.args.get("exam_id", type=int)
    report_view = request.args.get("view", "marks")
    if report_view not in {"marks", "points"}:
        abort(400)
    validate_report_exam(exam_id)
    report = build_student_report(student_id, exam_id)
    return render_template(
        "report_card_cbe.html",
        report=report,
        print_mode=True,
        report_view=report_view,
        selected_exam_id=exam_id,
    )


@app.route("/report/<int:student_id>/download")
@role_required("admin", "exam_officer", "teacher")
def download_report(student_id: int):
    ensure_teacher_student_access(student_id)
    exam_id = request.args.get("exam_id", type=int)
    report_view = request.args.get("view", "marks")
    if report_view not in {"marks", "points"}:
        abort(400)
    validate_report_exam(exam_id)
    report = build_student_report(student_id, exam_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Learner Result"
    sheet["A1"] = get_current_school()["name"]
    sheet["A2"] = f"{report['student']['full_name']} · {report['student']['grade']}"
    sheet["A3"] = report["exam_name"]
    if report_view == "points":
        sheet.append([])
        sheet.append(["Learning Area", "CBE Grade", "Points"])
        for item in report["subject_summary"]:
            sheet.append([item["name"], item["grade"]["grade"], item["points"]])
    else:
        sheet.append([])
        sheet.append(["Learning Area", "Marks", "Out of", "Percentage", "CBE Grade", "Points", "Exam"])
        for item in report["mark_entries"]:
            sheet.append(
                [
                    item["name"],
                    item["score"],
                    item["out_of"],
                    item["percentage"],
                    item["grade"],
                    item["points"],
                    item["exam_name"],
                ]
            )
    sheet.column_dimensions["A"].width = 30
    for column in ("B", "C", "D", "E", "F", "G"):
        sheet.column_dimensions[column].width = 18
    workbook_stream = BytesIO()
    workbook.save(workbook_stream)
    workbook_stream = embed_school_logo(
        workbook_stream, get_current_school().get("logo_url")
    )
    return send_file(
        workbook_stream,
        as_attachment=True,
        download_name=secure_filename(
            f"{report['student']['full_name']}-{report_view}-result.xlsx"
        ),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/class-analysis")
@role_required("admin", "exam_officer", "teacher")
def class_analysis():
    school_id = get_active_school_id()
    with get_db() as conn:
        if session.get("teacher_role") == "teacher":
            student_rows = conn.execute(
                """
                SELECT DISTINCT st.* FROM students st
                JOIN teacher_learning_areas a
                  ON a.school_id = st.school_id AND a.grade = st.grade
                WHERE st.school_id = ? AND a.teacher_id = ?
                ORDER BY st.grade, st.class_name, st.full_name
                """,
                (school_id, session.get("teacher_id")),
            ).fetchall()
            subjects = conn.execute(
                """
                SELECT DISTINCT s.* FROM subjects s
                JOIN teacher_learning_areas a ON a.subject_id = s.id
                WHERE a.school_id = ? AND a.teacher_id = ?
                ORDER BY s.name
                """,
                (school_id, session.get("teacher_id")),
            ).fetchall()
        else:
            student_rows = conn.execute("SELECT * FROM students WHERE school_id = ? ORDER BY grade, class_name, full_name", (school_id,)).fetchall()
            subjects = conn.execute("SELECT * FROM subjects WHERE school_id = ? ORDER BY name", (school_id,)).fetchall()

    analyses: List[Dict[str, Any]] = []
    grouped: Dict[str, List[Dict[str, Any]]] = {}
    for student in student_rows:
        group = f"{student['grade']} | {student['class_name']}"
        grouped.setdefault(group, []).append(student)

    for group_name, students in sorted(grouped.items()):
        students_reports = [build_student_report(student["id"]) for student in students]
        averages = []
        for subject in subjects:
            values = []
            for report in students_reports:
                for item in report["subject_summary"]:
                    if item["name"] == subject["name"]:
                        values.append(item["average"])
            if values:
                averages.append((subject["name"], round(sum(values) / len(values), 1)))
        analyses.append(
            {
                "group_name": group_name,
                "student_count": len(students_reports),
                "average": round(sum(report["overall_average"] for report in students_reports) / len(students_reports), 1) if students_reports else 0,
                "subject_averages": averages,
            }
        )

    return render_template("class_analysis.html", analyses=analyses)


@app.route("/settings", methods=["GET", "POST"])
@role_required("admin")
def settings():
    school_id = get_active_school_id()
    with get_db() as conn:
        current_settings = conn.execute("SELECT * FROM schools WHERE id = ?", (school_id,)).fetchone()

    if request.method == "POST":
        logo_url = request.form.get("logo_url")
        signature_url = request.form.get("signature_url")
        if "logo_file" in request.files:
            file = request.files["logo_file"]
            if file and file.filename:
                filename = secure_filename(file.filename)
                path = os.path.join(app.config["UPLOAD_FOLDER"], filename)
                file.save(path)
                logo_url = f"/static/uploads/{filename}"
        if "signature_file" in request.files:
            file = request.files["signature_file"]
            if file and file.filename:
                filename = secure_filename(file.filename)
                path = os.path.join(app.config["UPLOAD_FOLDER"], filename)
                file.save(path)
                signature_url = f"/static/uploads/{filename}"
        with get_db() as conn:
            conn.execute(
                "UPDATE schools SET name = ?, motto = ?, logo_url = ?, signature_url = ?, report_template = ? WHERE id = ?",
                (
                    request.form.get("name"),
                    request.form.get("motto"),
                    logo_url,
                    signature_url,
                    request.form.get("report_template"),
                    school_id,
                ),
            )
        flash("School branding updated.")
        return redirect(url_for("settings"))

    return render_template("settings.html", current_settings=current_settings)


if __name__ == "__main__":
    app.run(debug=True)
